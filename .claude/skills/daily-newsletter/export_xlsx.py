#!/usr/bin/env python3
"""Write data/markets.xlsx from data/markets.json: one sheet of headline figures, one per stock list.

Usage: python3 .claude/skills/daily-newsletter/export_xlsx.py
Needs openpyxl; if it is missing, says so and exits without failing the run.
"""
import json
import os
import sys

ROOT = os.path.normpath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
HEADLINE = [("usd_egp", "USD/EGP"), ("eur_egp", "EUR/EGP"), ("sar_egp", "SAR/EGP"), ("aed_egp", "AED/EGP"),
            ("kwd_egp", "KWD/EGP"), ("qar_egp", "QAR/EGP"), ("gbp_egp", "GBP/EGP"), ("dxy", "US dollar index"),
            ("gold21_egp", "Gold 21K (EGP/g)"), ("gold24_egp", "Gold 24K (EGP/g)"), ("gold18_egp", "Gold 18K (EGP/g)"),
            ("goldpound_egp", "Gold pound (EGP)"), ("gold_usd", "Gold ($/oz)"), ("brent", "Brent ($/bbl)"),
            ("egx30", "EGX30"), ("sp500", "S&P 500"), ("nasdaq", "Nasdaq"), ("dow", "Dow Jones"), ("btc", "Bitcoin ($)"),
            ("cbe_usd_buy", "USD/EGP central bank buy"), ("cbe_usd_sell", "USD/EGP central bank sell"),
            ("cbe_deposit", "Deposit rate (%)"), ("cbe_lending", "Lending rate (%)"), ("tbill_91", "91-day T-bill (%)"),
            ("tbill_182", "182-day T-bill (%)"), ("tbill_273", "273-day T-bill (%)"), ("tbill_364", "364-day T-bill (%)"),
            ("egx_turnover", "EGX turnover (EGP bn)"), ("infl_headline", "Inflation, headline (%)"), ("infl_core", "Inflation, core (%)"), ("reserves_usd", "Net international reserves ($bn)")]


def main():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl is not installed; skipped data/markets.xlsx")
        return
    days = json.load(open(os.path.join(ROOT, "data", "markets.json"))).get("days", [])
    wb = Workbook()
    head_font, head_fill = Font(bold=True), PatternFill("solid", fgColor="F0EBE0")

    def sheet(ws, columns, value):
        ws.append(["Date"] + [label for _, label in columns] + ["Backfilled"])
        for d in days:
            ws.append([d["date"]] + [value(d, k) for k, _ in columns] + ["yes" if d.get("seed") else ""])
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        ws.freeze_panes = "B2"
        for i in range(1, len(columns) + 3):
            ws.column_dimensions[get_column_letter(i)].width = 12 if i == 1 else 16
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=len(columns) + 1):
            for c in row:
                c.number_format = "#,##0.00"

    sheet(wb.active, HEADLINE, lambda d, k: d.get(k))
    wb.active.title = "Headline figures"
    for name, group in (("Egypt stocks", "egx"), ("US stocks", "us"), ("Volume (shares)", "vol")):
        tickers = sorted({t for d in days for t in d.get(group, {})})
        sheet(wb.create_sheet(name), [(t, t) for t in tickers], lambda d, k, g=group: d.get(g, {}).get(k))
    out = os.path.join(ROOT, "data", "markets.xlsx")
    wb.save(out)
    print(f"Wrote {os.path.relpath(out, ROOT)} ({len(days)} days)")


if __name__ == "__main__":
    sys.exit(main())
